import io
import os
import sys
import types

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from engine.asset_model import Asset, normalize_asset_state
from engine.export_engine import export_results_xlsx
from engine.governance import build_run_manifest


def _make_asset(**overrides):
    defaults = dict(
        id="A1",
        name="Asset One",
        lat=51.5,
        lon=-0.1,
        asset_type="residential_masonry",
        replacement_value=10_000_000,
        construction_material="masonry",
        year_built=2000,
        stories=2,
        basement=False,
        roof_type="gable",
        first_floor_height_m=0.0,
        terrain_elevation_asl_m=10.0,
        floor_area_m2=200.0,
        region="GBR",
    )
    defaults.update(overrides)
    return Asset(**defaults)


def _sample_annual_df() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "asset_id": "A1",
                "asset_name": "Asset One",
                "scenario_id": "current_policies",
                "year": 2025,
                "hazard": "flood",
                "ead": 1000.0,
                "pv": 1000.0,
                "data_source": "fallback_baseline",
            },
            {
                "asset_id": "A1",
                "asset_name": "Asset One",
                "scenario_id": "current_policies",
                "year": 2050,
                "hazard": "flood",
                "ead": 1500.0,
                "pv": 620.0,
                "data_source": "fallback_baseline",
            },
            {
                "asset_id": "A1",
                "asset_name": "Asset One",
                "scenario_id": "current_policies",
                "year": 2025,
                "hazard": "heat",
                "ead": 500.0,
                "pv": 500.0,
                "data_source": "isimip3b",
            },
            {
                "asset_id": "A2",
                "asset_name": "Asset Two",
                "scenario_id": "current_policies",
                "year": 2025,
                "hazard": "wildfire",
                "ead": 250.0,
                "pv": 250.0,
                "data_source": "manual_override",
            },
        ]
    )


def _sample_override_records() -> list[dict]:
    return [
        {
            "Asset ID": "A2",
            "Asset Name": "Asset Two",
            "Hazard": "wildfire",
            "Override basis": "Site survey / engineering assessment",
            "Source / justification": "Analyst override for site-specific screening.",
            "Prepared by": "tester",
            "Prepared at (UTC)": "2026-03-23T12:00:00Z",
            "Replaces source": "isimip3b",
            "Return periods": "10, 50, 100",
            "Override intensities": "0.1000, 0.2000, 0.3000",
        }
    ]


def _sample_manifest() -> dict:
    return build_run_manifest(
        annual_damages_df=_sample_annual_df(),
        selected_scenarios=["current_policies"],
        years=[2025, 2050],
        currency_code="USD",
        currency_symbol="$",
        discount_rate=0.035,
        fetch_profile="balanced",
        override_records=_sample_override_records(),
        fetch_failures=["Asset One: Aqueduct timeout"],
        reused_assets=1,
        refreshed_assets=1,
        zone_overrides={"A1": "EUR"},
    )


def test_normalize_asset_state_rewrites_dict_assets_in_place():
    asset = _make_asset()
    state = {"assets": [asset.to_dict(), asset]}

    normalized = normalize_asset_state(state)

    assert all(isinstance(item, Asset) for item in normalized)
    assert all(isinstance(item, Asset) for item in state["assets"])
    assert normalized[0].id == asset.id
    assert normalized[0].replacement_value == asset.replacement_value


def test_build_run_manifest_records_sources_fallbacks_and_overrides():
    manifest = _sample_manifest()

    assert manifest["Fetch profile"] == "balanced"
    assert manifest["Provider failure count"] == 1
    assert manifest["Fallback asset-hazard pairs"] == 1
    assert manifest["Manual override records"] == 1
    assert manifest["Zone override mode"] == "Preview only"

    provenance = manifest["Hazard provenance"]
    assert len(provenance) == 3, "Manifest should deduplicate repeated year rows to one asset-hazard lineage row"
    fallback_row = next(row for row in provenance if row["Hazard"] == "flood")
    override_row = next(row for row in provenance if row["Hazard"] == "wildfire")

    assert fallback_row["Fallback used"] == "Yes"
    assert override_row["Manual override"] == "Yes"
    assert override_row["Override replaces source"] == "isimip3b"
    assert any(event["Event type"] == "Provider refresh failure" for event in manifest["Provider events"])
    assert any(event["Event type"] == "Preview-only zone overrides" for event in manifest["Provider events"])


def test_results_export_includes_manifest_sheets():
    manifest = _sample_manifest()
    workbook_bytes = export_results_xlsx(
        asset_results_df=pd.DataFrame([{"Asset": "Asset One", "Scenario": "Current Policies", "EAD": 1000.0}]),
        annual_damages_df=_sample_annual_df(),
        portfolio_summary={"Assets": 2},
        scenarios=["current_policies"],
        metadata={"currency_symbol": "$"},
        run_manifest=manifest,
        override_records=_sample_override_records(),
    )

    workbook = load_workbook(io.BytesIO(workbook_bytes))

    assert "Run Metadata" in workbook.sheetnames
    assert "Hazard Provenance" in workbook.sheetnames
    assert "GEV Status" in workbook.sheetnames
    assert "Provider Events" in workbook.sheetnames
    assert "Manual Overrides" in workbook.sheetnames
    assert workbook["Hazard Provenance"]["A2"].value == "A1"
    assert workbook["Hazard Provenance"]["B2"].value == "Asset One"


def test_aqueduct_bws_cache_reuses_identical_requests():
    import engine.water_stress as water_stress

    calls = []

    class _Response:
        status_code = 200

        @staticmethod
        def json():
            return {"data": [{"indicator": "bws", "value": 2.5}]}

    def _fake_get(url, params=None, timeout=None):
        calls.append((url, params["geometry"], timeout))
        return _Response()

    original_requests = sys.modules.get("requests")
    water_stress._fetch_aqueduct_bws_cached.cache_clear()
    sys.modules["requests"] = types.SimpleNamespace(get=_fake_get)
    try:
        assert water_stress.fetch_aqueduct_bws(10.12344, 20.98764) == 2.5
        assert water_stress.fetch_aqueduct_bws(10.12341, 20.98761) == 2.5
    finally:
        water_stress._fetch_aqueduct_bws_cached.cache_clear()
        if original_requests is None:
            sys.modules.pop("requests", None)
        else:
            sys.modules["requests"] = original_requests

    assert len(calls) == 1, f"Expected a single cached Aqueduct request, got {len(calls)}"


def test_balanced_fetch_note_discloses_wildfire_fallback():
    for rel_path in ("pages/03_Hazards.py", "pages/04_Results.py"):
        path = os.path.join(os.path.dirname(__file__), "..", rel_path)
        with open(path, encoding="utf-8") as handle:
            source = handle.read()
        assert "Wildfire remains on the screening fallback baseline unless Full is selected." in source


def test_methodology_copy_no_longer_claims_weighted_tcfd_ready_dcf():
    path = os.path.join(os.path.dirname(__file__), "..", "pages", "00_Methodology.py")
    with open(path, encoding="utf-8") as handle:
        source = handle.read()

    assert "Scenario-weighted impairment %. TCFD-ready." not in source
    assert "weighted_npv" not in source
    assert "Scenario-specific impairment screening" in source


def test_results_and_audit_copy_explain_deferred_gev_bands():
    for rel_path in ("pages/04_Results.py", "pages/08_Audit.py"):
        path = os.path.join(os.path.dirname(__file__), "..", rel_path)
        with open(path, encoding="utf-8") as handle:
            source = handle.read()
        assert "Generate conditional GEV bands" in source
        assert "were not generated in the standard run" in source
