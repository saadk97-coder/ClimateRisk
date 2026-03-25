import io
import os
import shutil
import sys
import tempfile

import numpy as np
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from engine.adaptation_engine import calc_adaptation_bundle_npv, calc_adaptation_npv
from engine.asset_model import Asset
from engine.correlation import hazard_pair_correlation
from engine.export_engine import export_adaptation_xlsx, export_results_xlsx
from engine.governance import build_run_manifest
import engine.hazard_fetcher as hazard_fetcher
from engine.hazard_math import compute_effective_intensities, display_return_period_mask
from engine.uncertainty import fit_gev_central, fit_gev_return_levels

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))


def _make_asset(**overrides):
    defaults = dict(
        id="A1",
        name="Asset One",
        lat=51.5,
        lon=-0.1,
        asset_type="residential_masonry",
        replacement_value=10_000_000,
        construction_material="masonry",
        year_built=2000,
        stories=2,
        basement=False,
        roof_type="gable",
        first_floor_height_m=0.0,
        terrain_elevation_asl_m=10.0,
        floor_area_m2=200.0,
        region="GBR",
    )
    defaults.update(overrides)
    return Asset(**defaults)


def test_display_return_period_mask_hides_rp1000_from_standard_view():
    rp = np.array([10, 50, 100, 250, 500, 1000], dtype=float)
    mask = display_return_period_mask(rp)

    assert mask.tolist() == [True, True, True, True, True, False]


def test_gev_parameter_band_widens_in_tail():
    from scipy.stats import genextreme

    rng = np.random.default_rng(7)
    annual_maxima = genextreme.rvs(-0.1, loc=40.0, scale=8.0, size=45, random_state=rng)
    summary = fit_gev_return_levels(annual_maxima, np.array([50.0, 500.0]))

    assert summary is not None
    widths = np.asarray(summary["upper"]) - np.asarray(summary["lower"])
    assert widths[1] > widths[0], f"Expected wider tail band, got widths {widths}"


def test_gev_central_matches_legacy_central_curve():
    from scipy.stats import genextreme

    rng = np.random.default_rng(19)
    annual_maxima = genextreme.rvs(-0.08, loc=35.0, scale=6.5, size=40, random_state=rng)
    return_periods = np.array([10.0, 50.0, 100.0, 250.0, 500.0])

    central = fit_gev_central(annual_maxima, return_periods)
    legacy = fit_gev_return_levels(annual_maxima, return_periods)

    assert central is not None
    assert legacy is not None
    assert np.allclose(np.asarray(central["central"]), np.asarray(legacy["central"]))


def test_default_hazard_fetch_marks_deferred_without_eager_uncertainty(monkeypatch):
    original_cache_dir = hazard_fetcher._DISK_CACHE_DIR
    original_fetch_impl = hazard_fetcher._fetch_hazard_intensities_impl
    local_tmp = tempfile.mkdtemp(prefix="hazard-cache-", dir=PROJECT_ROOT)
    hazard_fetcher._DISK_CACHE_DIR = local_tmp
    hazard_fetcher._fetch_hazard_intensities_cached.cache_clear()

    gev_basis = {
        "version": 1,
        "return_periods": [10, 50, 100, 250, 500],
        "transform": {"kind": "identity"},
        "gcm_records": [
            {"gcm": "test", "annual_maxima": [25.0, 27.0, 29.5, 31.0, 28.0, 30.0, 32.0, 33.5, 34.0, 35.0, 36.0, 34.5]}
        ],
    }

    def _fake_fetch(*args, **kwargs):
        return (
            np.array([10, 50, 100, 250, 500], dtype=float),
            np.array([0.6, 1.2, 1.8, 2.4, 3.0], dtype=float),
            "isimip3b",
            {"gev_basis": gev_basis},
        )

    monkeypatch.setattr(hazard_fetcher, "_fetch_hazard_intensities_impl", _fake_fetch)
    try:
        _, _, source, diagnostics = hazard_fetcher.fetch_hazard_intensities(
            51.5,
            -0.1,
            "flood",
            "GBR",
            include_diagnostics=True,
        )
    finally:
        hazard_fetcher._DISK_CACHE_DIR = original_cache_dir
        hazard_fetcher._fetch_hazard_intensities_impl = original_fetch_impl
        hazard_fetcher._fetch_hazard_intensities_cached.cache_clear()
        shutil.rmtree(local_tmp, ignore_errors=True)

    assert source == "isimip3b"
    assert diagnostics.get("uncertainty_status") == "deferred"
    assert "uncertainty" not in diagnostics
    assert diagnostics.get("gev_basis")


def test_generate_conditional_gev_bands_uses_cached_basis_without_refetch(monkeypatch):
    original_cache_dir = hazard_fetcher._DISK_CACHE_DIR
    original_fetch_impl = hazard_fetcher._fetch_hazard_intensities_impl
    local_tmp = tempfile.mkdtemp(prefix="hazard-cache-", dir=PROJECT_ROOT)
    hazard_fetcher._DISK_CACHE_DIR = local_tmp
    hazard_fetcher._fetch_hazard_intensities_cached.cache_clear()

    gev_basis = {
        "version": 1,
        "return_periods": [10, 50, 100, 250, 500],
        "transform": {"kind": "identity"},
        "gcm_records": [
            {"gcm": "test", "annual_maxima": [28.0, 30.0, 32.0, 31.0, 33.0, 35.0, 36.0, 37.5, 38.0, 39.0, 40.0, 41.0]}
        ],
    }

    def _fake_fetch(*args, **kwargs):
        return (
            np.array([10, 50, 100, 250, 500], dtype=float),
            np.array([0.8, 1.4, 2.0, 2.8, 3.6], dtype=float),
            "isimip3b",
            {"gev_basis": gev_basis},
        )

    monkeypatch.setattr(hazard_fetcher, "_fetch_hazard_intensities_impl", _fake_fetch)
    try:
        rp, intensities, source, diagnostics = hazard_fetcher.fetch_hazard_intensities(
            40.7,
            -74.0,
            "heat",
            "USA",
            include_diagnostics=True,
        )

        def _should_not_refetch(*args, **kwargs):
            raise AssertionError("generate_conditional_gev_bands should use cached basis without refetching")

        monkeypatch.setattr(hazard_fetcher, "_fetch_hazard_intensities_impl", _should_not_refetch)
        entry = hazard_fetcher.generate_conditional_gev_bands(
            40.7,
            -74.0,
            "heat",
            "USA",
        )
    finally:
        hazard_fetcher._DISK_CACHE_DIR = original_cache_dir
        hazard_fetcher._fetch_hazard_intensities_impl = original_fetch_impl
        hazard_fetcher._fetch_hazard_intensities_cached.cache_clear()
        shutil.rmtree(local_tmp, ignore_errors=True)

    assert source == "isimip3b"
    assert diagnostics.get("uncertainty_status") == "deferred"
    assert entry["uncertainty_status"] == "generated"
    assert entry["uncertainty"]["type"] == "gev_parameter_uncertainty"
    assert np.allclose(np.asarray(entry["return_periods"], dtype=float), rp)
    assert np.allclose(np.asarray(entry["intensities"], dtype=float), intensities)


def test_results_copy_uses_scenario_range_not_uncertainty():
    results_path = os.path.join(os.path.dirname(__file__), "..", "pages", "04_Results.py")
    methodology_path = os.path.join(os.path.dirname(__file__), "..", "pages", "00_Methodology.py")

    with open(results_path, encoding="utf-8") as handle:
        results_source = handle.read()
    with open(methodology_path, encoding="utf-8") as handle:
        methodology_source = handle.read()

    assert "scenario range across selected pathways, not a statistical uncertainty interval" in results_source
    assert "Scenario comparison shows scenario range across selected pathways" in methodology_source


def test_hazard_specific_distance_decay_is_monotonic_and_heat_is_longer_range():
    flood_near = hazard_pair_correlation("flood", 25.0)
    flood_far = hazard_pair_correlation("flood", 1000.0)
    heat_far = hazard_pair_correlation("heat", 1000.0)

    assert flood_near > flood_far
    assert heat_far > flood_far


def test_coastal_effective_depth_subtracts_terrain_outside_multiplier():
    asset = _make_asset(first_floor_height_m=0.3, terrain_elevation_asl_m=2.0)
    base = np.array([0.8], dtype=float)

    effective, ctx = compute_effective_intensities(
        "coastal_flood",
        base,
        1.25,
        asset,
        scenario_id="current_policies",
        year=2050,
        region_zone="EUR",
    )

    expected = max(0.0, 0.8 * 1.25 + ctx.slr_additive_m - 2.0 - 0.3)
    assert abs(float(effective[0]) - expected) < 1e-9


def test_adaptation_bundle_uses_residual_risk_not_sum_of_standalone():
    annual_eads = {
        2025: 100_000.0,
        2026: 110_000.0,
        2027: 120_000.0,
        2028: 130_000.0,
        2029: 140_000.0,
        2030: 150_000.0,
    }
    hazard_streams = {"flood": dict(annual_eads)}
    measure_ids = ["flood_barrier_temporary", "flood_waterproofing"]

    standalone_one = calc_adaptation_npv(
        measure_ids[0],
        "A1",
        10_000_000.0,
        annual_eads,
        implementation_year=2026,
    )
    standalone_two = calc_adaptation_npv(
        measure_ids[1],
        "A1",
        10_000_000.0,
        annual_eads,
        implementation_year=2026,
    )
    bundle = calc_adaptation_bundle_npv(
        measure_ids,
        "A1",
        10_000_000.0,
        annual_eads,
        hazard_streams,
        implementation_year=2026,
    )

    assert bundle.bundled_npv_avoided_damages <= (
        standalone_one.npv_avoided_damages + standalone_two.npv_avoided_damages + 1e-6
    )
    second_row = bundle.measure_rows[1]
    assert second_row["incremental_npv_avoided"] < second_row["standalone_npv_avoided"]


def test_adaptation_export_includes_bundle_sheets():
    workbook_bytes = export_adaptation_xlsx(
        adaptation_df=pd.DataFrame([{"measure": "temporary barrier", "cbr": 1.2}]),
        frontier_df=pd.DataFrame([{"measure": "temporary barrier", "cumulative_capex": 100.0}]),
        bundle_df=pd.DataFrame([{"Sequence": 1, "Measure": "temporary barrier"}]),
        bundle_cashflows_df=pd.DataFrame([{"year": 2026, "avoided_damage": 10.0}]),
    )

    workbook = load_workbook(io.BytesIO(workbook_bytes))
    assert "Bundled Sequence" in workbook.sheetnames
    assert "Bundled Cashflows" in workbook.sheetnames


def test_run_manifest_captures_gev_diagnostics_rows():
    annual_df = pd.DataFrame(
        [
            {
                "asset_id": "A1",
                "asset_name": "Asset One",
                "scenario_id": "current_policies",
                "year": 2050,
                "hazard": "flood",
                "ead": 1200.0,
                "pv": 600.0,
                "data_source": "isimip3b",
            }
        ]
    )
    hazard_data_all = {
        "A1": {
            "flood": {
                "source": "isimip3b",
                "uncertainty": {
                    "type": "gev_parameter_uncertainty",
                    "return_periods": [50, 500],
                    "central": [1.2, 2.6],
                    "lower": [1.0, 2.0],
                    "upper": [1.5, 3.5],
                    "bootstrap_draws": 200,
                    "sample_years": 35,
                    "band_label": "P10-P90 parameter band",
                    "method": "GEV MLE with parametric bootstrap refits",
                    "limitation": "Conditional parameter uncertainty only.",
                },
            }
        }
    }

    manifest = build_run_manifest(
        annual_damages_df=annual_df,
        selected_scenarios=["current_policies"],
        years=[2050],
        currency_code="USD",
        currency_symbol="$",
        discount_rate=0.035,
        fetch_profile="balanced",
        hazard_data_all=hazard_data_all,
    )

    rows = manifest["GEV diagnostics"]
    assert len(rows) == 2
    assert {row["Hazard"] for row in rows} == {"flood"}
    assert {row["Bootstrap draws"] for row in rows} == {200}


def test_standard_results_export_omits_gev_diagnostics_but_records_status():
    annual_df = pd.DataFrame(
        [
            {
                "asset_id": "A1",
                "asset_name": "Asset One",
                "scenario_id": "current_policies",
                "year": 2050,
                "hazard": "flood",
                "ead": 900.0,
                "pv": 450.0,
                "data_source": "isimip3b",
            }
        ]
    )
    manifest = build_run_manifest(
        annual_damages_df=annual_df,
        selected_scenarios=["current_policies"],
        years=[2050],
        currency_code="USD",
        currency_symbol="$",
        discount_rate=0.035,
        fetch_profile="balanced",
        hazard_data_all={
            "A1": {
                "flood": {
                    "source": "isimip3b",
                    "uncertainty_status": "deferred",
                    "uncertainty_detail": "Conditional GEV parameter bands are available on demand and were not generated in the standard run.",
                    "gev_basis": {
                        "version": 1,
                        "return_periods": [10, 50, 100, 250, 500],
                        "transform": {"kind": "identity"},
                        "gcm_records": [{"gcm": "test", "annual_maxima": [1.0] * 12}],
                    },
                }
            }
        },
    )

    workbook = load_workbook(
        io.BytesIO(
            export_results_xlsx(
                asset_results_df=pd.DataFrame([{"Asset": "Asset One", "EAD": 900.0}]),
                annual_damages_df=annual_df,
                portfolio_summary={"Assets": 1},
                scenarios=["current_policies"],
                metadata={"currency_symbol": "$"},
                run_manifest=manifest,
            )
        )
    )

    assert "GEV Status" in workbook.sheetnames
    assert "GEV Diagnostics" not in workbook.sheetnames
    assert workbook["GEV Status"]["F2"].value == "deferred"


def test_full_results_export_includes_gev_diagnostics_when_generated():
    annual_df = pd.DataFrame(
        [
            {
                "asset_id": "A1",
                "asset_name": "Asset One",
                "scenario_id": "current_policies",
                "year": 2050,
                "hazard": "flood",
                "ead": 900.0,
                "pv": 450.0,
                "data_source": "isimip3b",
            }
        ]
    )
    manifest = build_run_manifest(
        annual_damages_df=annual_df,
        selected_scenarios=["current_policies"],
        years=[2050],
        currency_code="USD",
        currency_symbol="$",
        discount_rate=0.035,
        fetch_profile="balanced",
        hazard_data_all={
            "A1": {
                "flood": {
                    "source": "isimip3b",
                    "uncertainty": {
                        "type": "gev_parameter_uncertainty",
                        "return_periods": [50, 500],
                        "central": [1.2, 2.1],
                        "lower": [1.0, 1.7],
                        "upper": [1.5, 2.8],
                        "bootstrap_draws": 300,
                        "sample_years": 30,
                        "band_label": "Conditional parameter band",
                        "method": "GEV bootstrap refits",
                        "limitation": "Conditional parameter uncertainty only.",
                    },
                }
            }
        },
    )

    workbook = load_workbook(
        io.BytesIO(
            export_results_xlsx(
                asset_results_df=pd.DataFrame([{"Asset": "Asset One", "EAD": 900.0}]),
                annual_damages_df=annual_df,
                portfolio_summary={"Assets": 1},
                scenarios=["current_policies"],
                metadata={"currency_symbol": "$"},
                run_manifest=manifest,
            )
        )
    )

    assert "GEV Status" in workbook.sheetnames
    assert "GEV Diagnostics" in workbook.sheetnames
    assert workbook["GEV Status"]["F2"].value == "generated"
